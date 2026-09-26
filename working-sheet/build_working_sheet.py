"""Bangun working sheet Excel "Slot Map Weekly Meeting" dari tab Reguler Class Availability.

Pemakaian:
    python build_working_sheet.py <feed.csv> <output.xlsx> [--week "W40 · 28 Sep – 4 Okt 2026"] [--asof 2026-09-26]

<feed.csv> = CSV dari Sheet "Feed Slot Map (dari Exboard)" (IMPORTRANGE tab Reguler Class Availability).
"""
import argparse
import csv
import re
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CENTERS = ["KLM", "KWC", "PML", "TMP", "BTU", "HIB"]
CENTER_NAME = {"KLM": "Kalimalang", "KWC": "Karawaci", "PML": "Pamulang", "TMP": "TMP", "BTU": "BTU", "HIB": "HIB"}
DAYS = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
AGES = ["Baby", "Tiny", "Little", "Kids"]
AGE_RANGE = {"Baby": "6–17 bulan", "Tiny": "18–36 bulan", "Little": "3–5 tahun", "Kids": "5–8+ tahun"}
SESSION_GAP_MIN = 45  # dua sesi dianggap bentrok bila jaraknya < 45 menit
DECISIONS = ["Jual minggu ini", "Buka kelas baru", "Ganti age group", "Tahan", "Tutup slot"]

SLOT_RE = re.compile(r"^(.*?)-\s*Lane\s*(\d+)\s*-\s*([A-Za-z]+)\s*-\s*(\d{1,2})[:.](\d{2})")
F = {"age": 0, "day": 1, "time": 2, "slot": 3, "cap": 4, "used": 5, "avail": 6, "coach": 7}

# ---------- style ----------
FONT = "Arial"
C_INK, C_MUTED = "1F2D33", "6B7C82"
C_ACCENT, C_ACCENT_SOFT = "0B6475", "DCEDF0"
C_INPUT = "FFF4B8"
STATUS_FILL = {
    "Belum terisi": "F8C9C0",      # kelas ada, 0 murid  -> fokus diskusi
    "Slot kosong": "FFE3A6",       # kapasitas ada, belum ada kelas
    "Sebagian": "E3F1F4",
    "Penuh": "BFE3C8",
    "Lebih kapasitas": "E2D4F5",
    "Tutup": "E4E4E4",
}
AGE_FILL = {"Baby": "F7DBE6", "Tiny": "FDE7C4", "Little": "D6EFDC", "Kids": "DBE4FB"}
thin = Side(style="thin", color="D5DDDF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, size=10, color=C_INK, italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


# ---------- parse ----------
def cell_num(s):
    if s is None:
        return None
    s = re.sub(r"[\s  ]", "", str(s))
    if not s or s.startswith("#"):
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", s)
    if m:
        return (date(int(m[1]), int(m[2]), int(m[3])) - date(1899, 12, 30)).days
    try:
        n = float(s.replace(",", "."))
        return int(n) if n == int(n) else n
    except ValueError:
        return None


def clean(s):
    return re.sub(r"[  ]", " ", s or "").strip()


def title_case(s):
    return re.sub(r"(^|[\s-])(\w)", lambda m: m[1] + m[2].upper(), s.lower())


def norm_age(s):
    s = clean(s)
    if not s or re.fullmatch(r"[-–—.]+", s):
        return ""
    for a in AGES:
        if a.lower() == s.lower():
            return a
    return title_case(s)


def norm_coach(s):
    s = re.sub(r"\s+", " ", clean(s))
    if not s or s == "-" or re.fullmatch(r"(not )?available", s, re.I):
        return ""
    return title_case(s)


def parse_feed(path):
    """Sama dengan parser dashboard: deteksi pergeseran kolom dari posisi kolom Slot."""
    slots = []
    with open(path, encoding="utf-8") as fh:
        for r in csv.reader(fh):
            center = clean(r[0] if r else "").upper()
            if not re.fullmatch(r"[A-Z]{2,5}", center) or center == "CENTER":
                continue
            j, m = -1, None
            for k in range(1, min(10, len(r))):
                mm = SLOT_RE.match(clean(r[k]))
                if mm:
                    j, m = k, mm
                    break
            if j < 4:
                continue
            shift = j - 4

            def get(f):
                pos = 1 + shift + f
                if pos <= 8:
                    return r[pos] if pos < len(r) else ""
                w = pos - 10
                return r[w] if w >= 1 else None

            day = next((d for d in DAYS if d.lower() == m[3].lower()), None)
            if not day:
                continue
            slots.append({
                "center": center, "lane": int(m[2]), "day": day,
                "time": f"{int(m[4]):02d}:{m[5]}",
                "age": norm_age(re.sub(r"-\s*$", "", m[1])) or norm_age(get(F["age"])),
                "cap": cell_num(get(F["cap"])), "used": cell_num(get(F["used"])),
                "coach": norm_coach(get(F["coach"])),
            })
    return slots


def status_of(s):
    cap, used, age, coach = s["cap"], s["used"], s["age"], s["coach"]
    if cap is None or used is None:
        return "Data tidak lengkap"
    if cap == 0 and used == 0:
        return "Tutup" if (age or coach) else "Tidak dibuka"
    if used > cap:
        return "Lebih kapasitas"
    if used == cap:
        return "Penuh"
    if used == 0:
        return "Belum terisi" if age else "Slot kosong"
    return "Sebagian"


def minutes(t):
    h, m = t.split(":")
    return int(h) * 60 + int(m)


# ---------- coach availability ----------
def coach_book(slots):
    """Per center: jadwal tiap coach (hari -> set menit), age group yang diajar."""
    book = defaultdict(lambda: defaultdict(lambda: {"busy": defaultdict(set), "ages": set(), "sessions": 0, "name": ""}))
    for s in slots:
        if not s["coach"]:
            continue
        c = book[s["center"]][s["coach"].lower()]
        c["name"] = s["coach"]
        c["busy"][s["day"]].add(minutes(s["time"]))
        if s["age"]:
            c["ages"].add(s["age"])
        c["sessions"] += 1
    return book


def available_coaches(book_center, day, time, exclude=""):
    """Coach yang bertugas di center itu pada hari yang sama, tidak sedang mengajar dalam ±45 menit."""
    t = minutes(time)
    out = []
    for key, c in book_center.items():
        if key == exclude.lower():
            continue
        busy = c["busy"].get(day)
        if not busy:
            continue
        if all(abs(t - b) >= SESSION_GAP_MIN for b in busy):
            out.append(c)
    order = {a: i for i, a in enumerate(AGES)}
    out.sort(key=lambda c: (-c["sessions"], c["name"]))
    return out, order


def ages_str(ages):
    order = {a: i for i, a in enumerate(AGES)}
    return ", ".join(sorted(ages, key=lambda a: (order.get(a, 99), a)))


# ---------- workbook ----------
def build(slots, out_path, week_label, asof):
    wb = Workbook()
    book = coach_book(slots)
    centers = [c for c in CENTERS if any(s["center"] == c for s in slots)]

    # ===== Master Data =====
    md = wb.active
    md.title = "Master Data"
    headers = ["Key", "Center", "Lane", "Hari", "Jam", "Age Group", "Kapasitas", "Terisi",
               "Sisa Kursi", "Utilisasi", "Coach", "Status", "Label Slot Map"]
    widths = [22, 8, 6, 9, 7, 11, 10, 8, 10, 10, 14, 16, 30]
    md.append(headers)
    for i, w in enumerate(widths, 1):
        md.column_dimensions[get_column_letter(i)].width = w
    for c in md[1]:
        c.font = font(True, 10, "FFFFFF")
        c.fill = fill(C_ACCENT)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    day_idx = {d: i for i, d in enumerate(DAYS)}
    slots_sorted = sorted(slots, key=lambda s: (CENTERS.index(s["center"]) if s["center"] in CENTERS else 99,
                                                s["lane"], day_idx[s["day"]], s["time"]))
    for i, s in enumerate(slots_sorted, start=2):
        md.append([
            f'=B{i}&"|"&C{i}&"|"&D{i}&"|"&E{i}', s["center"], s["lane"], s["day"], s["time"],
            s["age"] or None, s["cap"], s["used"],
            f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),MAX(G{i}-H{i},0),"")',
            f'=IF(AND(ISNUMBER(G{i}),ISNUMBER(H{i})),IF(G{i}>0,H{i}/G{i},""),"")',
            s["coach"] or None,
            (f'=IF(OR(G{i}="",H{i}=""),"Data tidak lengkap",IF(AND(G{i}=0,H{i}=0),IF(AND(F{i}="",K{i}=""),"Tidak dibuka","Tutup"),'
             f'IF(H{i}>G{i},"Lebih kapasitas",IF(H{i}=G{i},"Penuh",IF(H{i}=0,IF(F{i}="","Slot kosong","Belum terisi"),"Sebagian")))))'),
            (f'=IF(L{i}="Tidak dibuka","–",IF(F{i}="","(tanpa age group)",F{i})&"  "&IF(H{i}="","?",H{i})&"/"&IF(G{i}="","?",G{i})'
             f'&CHAR(10)&IF(K{i}="","coach: –","coach: "&K{i}))'),
        ])
        md.cell(i, 10).number_format = "0%"
    last_md = len(slots_sorted) + 1
    for row in md.iter_rows(min_row=2, max_row=last_md):
        for c in row:
            c.font = font(size=9)
    for col in (6, 7, 8, 11):   # kolom input dari sumber
        for r in range(2, last_md + 1):
            md.cell(r, col).font = font(size=9, color="0000FF")
    md.freeze_panes = "B2"
    md.auto_filter.ref = f"A1:M{last_md}"
    md.cell(1, 7).comment = Comment("Kapasitas, Terisi, Age Group dan Coach disalin dari Exboard › Reguler Class Availability. "
                                    "Kolom lain dihitung dengan rumus.", "Working sheet")
    M = "'Master Data'"
    rng = lambda col: f"{M}!${col}$2:${col}${last_md}"

    # ===== Ringkasan =====
    rk = wb.create_sheet("Ringkasan", 0)

    # ===== Cara Pakai =====
    cp = wb.create_sheet("Cara Pakai", 0)
    build_guide(cp, week_label, asof)

    build_summary(rk, centers, rng, week_label, asof)

    # ===== Per center =====
    for c in centers:
        ws = wb.create_sheet(c, wb.sheetnames.index("Master Data"))
        build_center(ws, c, [s for s in slots_sorted if s["center"] == c], book[c], rng, week_label, asof)

    wb.active = wb.sheetnames.index("Ringkasan")
    wb.save(out_path)


def build_guide(ws, week_label, asof):
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90
    ws.sheet_view.showGridLines = False
    r = 2
    ws.cell(r, 2, "Slot Map Weekly Meeting · Center Manager × Senior Coach").font = font(True, 16, C_ACCENT)
    r += 1
    ws.cell(r, 2, f"Data per {asof} · untuk meeting {week_label} · sumber: Exboard › Reguler Class Availability").font = font(size=10, color=C_MUTED)
    r += 2
    ws.cell(r, 2, "Alur meeting").font = font(True, 12)
    r += 1
    steps = [
        ("1. Buka tab center", "Satu tab per center (KLM, KWC, PML, TMP, BTU, HIB). Mulai dari ringkasan di atas: utilisasi, jumlah kelas belum terisi, slot kosong."),
        ("2. Lihat Slot Map", "Grid per lane: baris = jam, kolom = hari. Tiap sel berisi age group, terisi/kapasitas, dan coach. Merah = kelas sudah ada tapi 0 murid. Oranye = kapasitas ada tapi belum ada kelas."),
        ("3. Putuskan di Daftar Diskusi", "Di bawah slot map ada daftar semua slot merah & oranye. Isi kolom kuning: keputusan, age group yang dijual, coach PIC, target murid, catatan."),
        ("4. Cek coach", "Kolom 'Coach available' menyarankan coach yang bertugas hari itu dan tidak sedang mengajar di jam tersebut, beserta age group yang biasa mereka ajar. Tabel Coach di bagian bawah tab menunjukkan beban tiap coach per hari."),
        ("5. Kirim ke Student Advisor", "Filter kolom Keputusan = 'Jual minggu ini' atau 'Buka kelas baru'. Itulah daftar kelas yang dijual SA minggu itu."),
    ]
    for a, b in steps:
        ws.cell(r, 2, a).font = font(True)
        ws.cell(r, 3, b).font = font()
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 2).alignment = Alignment(vertical="top")
        r += 1
    r += 1
    ws.cell(r, 2, "Arti warna status").font = font(True, 12)
    r += 1
    meaning = [
        ("Belum terisi", "Kelas sudah punya age group, kapasitas > 0, tapi 0 murid. Prioritas untuk dijual."),
        ("Slot kosong", "Kapasitas > 0 tapi belum ada age group/kelas. Bisa dibuka jadi kelas baru."),
        ("Sebagian", "Sudah ada murid tapi belum penuh."),
        ("Penuh", "Terisi = kapasitas."),
        ("Lebih kapasitas", "Terisi melebihi kapasitas (termasuk slot kapasitas 0 yang masih ada murid). Perlu dicek."),
        ("Tutup", "Kapasitas 0 dan tidak ada murid, tapi slot masih punya age group/coach."),
    ]
    for st, txt in meaning:
        ws.cell(r, 2, st).fill = fill(STATUS_FILL[st])
        ws.cell(r, 2).font = font(True)
        ws.cell(r, 2).border = BORDER
        ws.cell(r, 3, txt).font = font()
        r += 1
    ws.cell(r, 2, "Tidak dibuka").font = font(True, color=C_MUTED)
    ws.cell(r, 3, "Kapasitas 0, tanpa kelas. Ditampilkan '–'.").font = font()
    r += 2
    ws.cell(r, 2, "Sel yang diisi").font = font(True, 12)
    r += 1
    ws.cell(r, 2, "Kuning").fill = fill(C_INPUT)
    ws.cell(r, 2).font = font(True)
    ws.cell(r, 2).border = BORDER
    ws.cell(r, 3, "Hanya sel kuning yang diisi saat meeting (minggu, peserta, keputusan, age group dijual, coach PIC, target, catatan). Sel lain berisi rumus.").font = font()
    r += 1
    ws.cell(r, 2, "Teks biru").font = font(True, color="0000FF")
    ws.cell(r, 3, "Di tab Master Data: angka yang disalin dari Exboard (age group, kapasitas, terisi, coach).").font = font()
    r += 2
    ws.cell(r, 2, "Contoh pengisian").font = font(True, 12)
    r += 1
    ex_h = ["Hari · Jam · Lane", "Status", "Coach available", "Keputusan", "Age group dijual", "Coach PIC", "Target murid", "Catatan"]
    ex_v = ["Selasa · 09:45 · Lane 2", "Belum terisi", "Danti (Baby, Little)", "Jual minggu ini", "Tiny", "Danti", 3, "Tawarkan ke leads trial Tiny minggu lalu"]
    for i, (h, v) in enumerate(zip(ex_h, ex_v)):
        ws.cell(r + i, 2, h).font = font(True, color=C_MUTED)
        ws.cell(r + i, 3, v).font = font(italic=True)
        ws.cell(r + i, 3).alignment = Alignment(horizontal="left")
    r += len(ex_h) + 1
    ws.cell(r, 2, "Catatan").font = font(True, 12)
    r += 1
    notes = [
        "Coach available = coach yang punya jadwal di center itu pada hari yang sama dan tidak mengajar dalam rentang ±45 menit dari slot tersebut. "
        "Age group di dalam kurung = age group yang dia ajar minggu ini di center itu. Ini saran; cek ulang ke senior coach.",
        "Coach yang tidak punya jadwal sama sekali di hari itu tidak muncul sebagai available (bisa saja libur).",
        "Untuk memperbarui data minggu berikutnya: ganti isi tab Master Data (kolom biru) dengan data terbaru. Slot map, ringkasan dan status ikut menghitung ulang. "
        "Daftar Diskusi dan saran coach dibuat saat file digenerate, jadi generate ulang file-nya tiap minggu.",
    ]
    for n in notes:
        ws.cell(r, 3, n).font = font()
        ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 2, "•").font = font(True)
        ws.cell(r, 2).alignment = Alignment(horizontal="right", vertical="top")
        r += 1


def build_summary(ws, centers, rng, week_label, asof):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.cell(2, 2, "Ringkasan Utilisasi per Center").font = font(True, 16, C_ACCENT)
    ws.cell(3, 2, f"Data per {asof} · meeting {week_label}").font = font(size=10, color=C_MUTED)
    heads = ["Center", "Utilisasi", "Kursi terisi", "Kapasitas", "Belum terisi\n(kelas 0 murid)", "Slot kosong\n(belum ada kelas)",
             "Sebagian", "Penuh", "Lebih kapasitas", "Buka tab"]
    widths = [14, 11, 12, 11, 15, 16, 10, 9, 12, 11]
    for i, (h, w) in enumerate(zip(heads, widths)):
        c = ws.cell(5, 2 + i, h)
        c.font = font(True, 10, "FFFFFF")
        c.fill = fill(C_ACCENT)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(2 + i)].width = w
    ws.row_dimensions[5].height = 30
    for k, cen in enumerate(centers):
        r = 6 + k
        ws.cell(r, 2, cen).font = font(True)
        ws.cell(r, 3, f'=IFERROR(D{r}/E{r},"")').number_format = "0%"
        ws.cell(r, 4, f'=SUMIFS({rng("H")},{rng("B")},B{r},{rng("G")},">0")')
        ws.cell(r, 5, f'=SUMIFS({rng("G")},{rng("B")},B{r},{rng("G")},">0")')
        for j, st in enumerate(["Belum terisi", "Slot kosong", "Sebagian", "Penuh", "Lebih kapasitas"]):
            cell = ws.cell(r, 6 + j, f'=COUNTIFS({rng("B")},$B{r},{rng("L")},"{st}")')
            cell.fill = fill(STATUS_FILL[st]) if st in ("Belum terisi", "Slot kosong") else PatternFill()
        link = ws.cell(r, 11, f"→ {cen}")
        link.hyperlink = f"#'{cen}'!A1"
        link.font = Font(name=FONT, size=10, color=C_ACCENT, underline="single")
        for col in range(2, 12):
            ws.cell(r, col).border = BORDER
            if col != 11:
                ws.cell(r, col).font = font(col == 2 or col == 3)
            ws.cell(r, col).alignment = Alignment(horizontal="center" if col > 2 else "left")
    tr = 6 + len(centers)
    ws.cell(tr, 2, "Total").font = font(True)
    ws.cell(tr, 3, f'=IFERROR(D{tr}/E{tr},"")').number_format = "0%"
    for col in range(4, 11):
        L = get_column_letter(col)
        ws.cell(tr, col, f"=SUM({L}6:{L}{tr-1})")
    for col in range(2, 11):
        ws.cell(tr, col).font = font(True)
        ws.cell(tr, col).fill = fill(C_ACCENT_SOFT)
        ws.cell(tr, col).border = BORDER
        ws.cell(tr, col).alignment = Alignment(horizontal="center" if col > 2 else "left")
    ws.cell(tr + 2, 2, "Utilisasi = kursi terisi ÷ kapasitas, hanya slot dengan kapasitas > 0.").font = font(size=9, color=C_MUTED)
    ws.conditional_formatting.add(f"C6:C{tr-1}", ColorScaleRule(start_type="num", start_value=0, start_color="F8C9C0",
                                                                 mid_type="num", mid_value=0.5, mid_color="FFF4B8",
                                                                 end_type="num", end_value=1, end_color="BFE3C8"))


def build_center(ws, center, slots, book_c, rng, week_label, asof):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    widths = {"A": 8, "B": 19, "C": 19, "D": 19, "E": 19, "F": 19, "G": 19, "H": 19,
              "I": 40, "J": 18, "K": 16, "L": 14, "M": 11, "N": 34}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    HELP0 = 20  # kolom bantu (status) mulai kolom T, disembunyikan
    for i in range(HELP0, HELP0 + 7):
        ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ----- Header -----
    name = CENTER_NAME.get(center, center)
    ws.cell(1, 1, f"SLOT MAP · {center}" + (f" · {name}" if name != center else "")).font = font(True, 16, C_ACCENT)
    ws.merge_cells("A1:F1")
    ws.cell(2, 1, f"Data per {asof} · sumber: Exboard › Reguler Class Availability").font = font(size=9, color=C_MUTED)
    ws.merge_cells("A2:F2")
    ws["A3"] = "Center"
    ws["B3"] = center
    ws["A4"] = "Minggu"
    ws["B4"] = week_label
    ws["C3"] = "Center Manager"
    ws["C4"] = "Senior Coach"
    for a in ("A3", "A4", "C3", "C4"):
        ws[a].font = font(True, color=C_MUTED)
    ws["B3"].font = font(True, 12)
    for a in ("B4", "D3", "D4"):
        ws[a].fill = fill(C_INPUT)
        ws[a].border = BORDER
        ws[a].font = font()
    ws.merge_cells("D3:E3")
    ws.merge_cells("D4:E4")

    # KPI row
    kpis = [
        ("Utilisasi", f'=IFERROR(SUMIFS({rng("H")},{rng("B")},$B$3,{rng("G")},">0")/SUMIFS({rng("G")},{rng("B")},$B$3,{rng("G")},">0"),"")', "0%", C_ACCENT_SOFT),
        ("Kursi terisi", f'=SUMIFS({rng("H")},{rng("B")},$B$3,{rng("G")},">0")&" / "&SUMIFS({rng("G")},{rng("B")},$B$3,{rng("G")},">0")', "@", C_ACCENT_SOFT),
        ("Belum terisi", f'=COUNTIFS({rng("B")},$B$3,{rng("L")},"Belum terisi")', "0", STATUS_FILL["Belum terisi"]),
        ("Slot kosong", f'=COUNTIFS({rng("B")},$B$3,{rng("L")},"Slot kosong")', "0", STATUS_FILL["Slot kosong"]),
        ("Sebagian", f'=COUNTIFS({rng("B")},$B$3,{rng("L")},"Sebagian")', "0", STATUS_FILL["Sebagian"]),
        ("Penuh", f'=COUNTIFS({rng("B")},$B$3,{rng("L")},"Penuh")', "0", STATUS_FILL["Penuh"]),
        ("Lebih kapasitas", f'=COUNTIFS({rng("B")},$B$3,{rng("L")},"Lebih kapasitas")', "0", STATUS_FILL["Lebih kapasitas"]),
    ]
    for i, (lab, f_, nf, col) in enumerate(kpis):
        c = 2 + i
        ws.cell(6, c, lab).font = font(True, 9, C_INK)
        ws.cell(6, c).fill = fill(col)
        v = ws.cell(7, c, f_)
        v.font = font(True, 16)
        v.number_format = nf
        v.fill = fill(col)
        v.alignment = Alignment(horizontal="left")
        for rr in (6, 7):
            ws.cell(rr, c).border = BORDER
    ws.cell(6, 1, "Ringkasan").font = font(True, 9, C_MUTED)

    # Legend
    ws.cell(9, 1, "Warna").font = font(True, 9, C_MUTED)
    for i, st in enumerate(["Belum terisi", "Slot kosong", "Sebagian", "Penuh", "Lebih kapasitas", "Tutup"]):
        c = ws.cell(9, 2 + i, st)
        c.fill = fill(STATUS_FILL[st])
        c.font = font(True, 9)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
    ws.cell(10, 2, "Isi sel: age group  terisi/kapasitas, lalu coach. Merah & oranye = bahan diskusi (lihat Daftar Diskusi di bawah).").font = font(size=9, color=C_MUTED, italic=True)

    # Navigation placeholders (filled after we know row numbers)
    nav_row = 11

    # ----- Slot map per lane -----
    r = 13
    slotmap_row = r
    lanes = sorted({s["lane"] for s in slots})
    grid_ranges = []
    for lane in lanes:
        ls = [s for s in slots if s["lane"] == lane]
        times = sorted({s["time"] for s in ls})
        ws.cell(r, 1, f"LANE {lane}").font = font(True, 13, "FFFFFF")
        for c in range(1, 9):
            ws.cell(r, c).fill = fill(C_ACCENT)
        ws.cell(r, 3, f'="Utilisasi lane: "&IFERROR(TEXT(SUMIFS({rng("H")},{rng("B")},$B$3,{rng("C")},{lane},{rng("G")},">0")/SUMIFS({rng("G")},{rng("B")},$B$3,{rng("C")},{lane},{rng("G")},">0"),"0%"),"–")'
                         f'&"   ·   Belum terisi: "&COUNTIFS({rng("B")},$B$3,{rng("C")},{lane},{rng("L")},"Belum terisi")'
                         f'&"   ·   Slot kosong: "&COUNTIFS({rng("B")},$B$3,{rng("C")},{lane},{rng("L")},"Slot kosong")').font = font(True, 10, "FFFFFF")
        r += 1
        ws.cell(r, 1, "Jam").font = font(True, 9, C_MUTED)
        for i, d in enumerate(DAYS):
            h = ws.cell(r, 2 + i, d)
            h.font = font(True, 10)
            h.fill = fill("F0F4F5")
            h.alignment = Alignment(horizontal="center")
            h.border = BORDER
        ws.cell(r, 1).border = BORDER
        head_row = r
        r += 1
        top = r
        for t in times:
            tc = ws.cell(r, 1, t)
            tc.font = font(True, 10)
            tc.alignment = Alignment(vertical="top", horizontal="center")
            tc.border = BORDER
            for i in range(7):
                col = 2 + i
                key = f'$B$3&"|{lane}|"&{get_column_letter(col)}${head_row}&"|"&$A{r}'
                c = ws.cell(r, col, f'=IFERROR(INDEX({rng("M")},MATCH({key},{rng("A")},0)),"")')
                c.font = font(size=9)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = BORDER
                hc = ws.cell(r, HELP0 + i, f'=IFERROR(INDEX({rng("L")},MATCH({key},{rng("A")},0)),"")')
                hc.font = font(size=8, color=C_MUTED)
            ws.row_dimensions[r].height = 30
            r += 1
        grid_ranges.append((top, r - 1))
        r += 1

    # Conditional formatting on grids via hidden helper status columns
    helper_first = get_column_letter(HELP0)
    for top, bot in grid_ranges:
        area = f"B{top}:H{bot}"
        for st, col in STATUS_FILL.items():
            ws.conditional_formatting.add(area, FormulaRule(formula=[f'{helper_first}{top}="{st}"'], fill=fill(col), stopIfTrue=True))
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'{helper_first}{top}="Tidak dibuka"'], font=Font(name=FONT, color="B0B8BA"), stopIfTrue=True))

    # ----- Daftar Diskusi -----
    r += 1
    diskusi_row = r
    ws.cell(r, 1, "DAFTAR DISKUSI · kelas yang belum terisi & slot kosong").font = font(True, 13, "FFFFFF")
    for c in range(1, 15):
        ws.cell(r, c).fill = fill("B03A2E")
    r += 1
    ws.cell(r, 1, "Isi kolom kuning saat meeting. Filter kolom Keputusan untuk daftar yang dijual Student Advisor minggu ini.").font = font(size=9, color=C_MUTED, italic=True)
    r += 1
    hdr = ["No", "Hari", "Jam", "Lane", "Status", "Age group terjadwal", "Kapasitas", "Coach terjadwal",
           "Coach available\n(age group yang diajar)", "Age group tersedia", "Keputusan", "Age group dijual", "Coach PIC", "Target murid", "Catatan"]
    # columns: A..O  (A=No, B=Hari, ... O=Catatan)
    for i, h in enumerate(hdr):
        c = ws.cell(r, 1 + i, h)
        c.font = font(True, 9, "FFFFFF")
        c.fill = fill(C_INK) if i < 10 else fill("7A5C00")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
    ws.column_dimensions["O"].width = 34
    ws.column_dimensions["N"].width = 11
    ws.column_dimensions["M"].width = 14
    ws.row_dimensions[r].height = 32
    list_head = r
    r += 1
    day_idx = {d: i for i, d in enumerate(DAYS)}
    todo = [s for s in slots if status_of(s) in ("Belum terisi", "Slot kosong")]
    todo.sort(key=lambda s: (day_idx[s["day"]], s["time"], s["lane"]))
    coach_names = sorted({c["name"] for c in book_c.values()})
    first_item = r
    for n, s in enumerate(todo, 1):
        key = f'$B$3&"|"&D{r}&"|"&B{r}&"|"&C{r}'
        avail, _ = available_coaches(book_c, s["day"], s["time"], exclude=s["coach"])
        avail_txt = "; ".join(f'{c["name"]} ({ages_str(c["ages"]) or "–"})' for c in avail[:6])
        if len(avail) > 6:
            avail_txt += f"; +{len(avail) - 6} lainnya"
        pool = set()
        if s["age"]:
            pool.add(s["age"])
        if s["coach"] and s["coach"].lower() in book_c:
            pool |= book_c[s["coach"].lower()]["ages"]
        for c in avail:
            pool |= c["ages"]
        row = [n, s["day"], s["time"], s["lane"],
               f'=IFERROR(INDEX({rng("L")},MATCH({key},{rng("A")},0)),"")',
               f'=IFERROR(INDEX({rng("F")},MATCH({key},{rng("A")},0))&"","")',
               f'=IFERROR(INDEX({rng("G")},MATCH({key},{rng("A")},0)),"")',
               f'=IFERROR(INDEX({rng("K")},MATCH({key},{rng("A")},0))&"","")',
               avail_txt or "Tidak ada coach lain yang bertugas & kosong",
               ages_str(pool) or "–",
               None, None, None, None, None]
        for i, v in enumerate(row):
            c = ws.cell(r, 1 + i, v)
            c.font = font(size=9, color=C_MUTED if (i == 8 and not avail_txt) else C_INK)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if i in (0, 2, 3, 6, 13) else "left")
            if i >= 10:
                c.fill = fill(C_INPUT)
        lines = max(1, -(-len(avail_txt) // 44))
        ws.row_dimensions[r].height = max(16, 12.5 * lines + 3)
        r += 1
    last_item = r - 1
    if todo:
        st_col = f"E{first_item}:E{last_item}"
        ws.conditional_formatting.add(st_col, FormulaRule(formula=[f'E{first_item}="Belum terisi"'], fill=fill(STATUS_FILL["Belum terisi"])))
        ws.conditional_formatting.add(st_col, FormulaRule(formula=[f'E{first_item}="Slot kosong"'], fill=fill(STATUS_FILL["Slot kosong"])))
        ws.conditional_formatting.add(f"A{first_item}:O{last_item}", FormulaRule(formula=[f'$K{first_item}="Jual minggu ini"'], font=Font(name=FONT, bold=True, color="0D6B3A")))
        ws.auto_filter.ref = f"A{list_head}:O{last_item}"
        dv1 = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"' + ",".join(AGES) + '"', allow_blank=True)
        dv4 = DataValidation(type="whole", operator="between", formula1="0", formula2="20", allow_blank=True,
                             error="Isi angka 0–20", errorTitle="Target murid")
        for dv in (dv1, dv2, dv4):
            ws.add_data_validation(dv)
        dv1.add(f"K{first_item}:K{last_item}")
        dv2.add(f"L{first_item}:L{last_item}")
        dv4.add(f"N{first_item}:N{last_item}")
    else:
        ws.cell(r, 1, "Tidak ada kelas 0 murid atau slot kosong di center ini.").font = font(italic=True, color=C_MUTED)
        r += 1

    # ----- Coach -----
    r += 2
    coach_row = r
    ws.cell(r, 1, "COACH · jadwal minggu ini di center ini").font = font(True, 13, "FFFFFF")
    for c in range(1, 15):
        ws.cell(r, c).fill = fill(C_ACCENT)
    r += 1
    ws.cell(r, 1, "Angka = jumlah sesi coach per hari (dihitung dari Master Data). Dipakai untuk memilih coach PIC.").font = font(size=9, color=C_MUTED, italic=True)
    r += 1
    ch = ["No", "Coach", "Age group diajar"] + DAYS + ["Total sesi", "Jam pertama–terakhir"]
    for i, h in enumerate(ch):
        c = ws.cell(r, 1 + i, h)
        c.font = font(True, 9, "FFFFFF")
        c.fill = fill(C_INK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    r += 1
    c_first = r
    coaches = sorted(book_c.values(), key=lambda c: (-c["sessions"], c["name"]))
    for n, cinfo in enumerate(coaches, 1):
        allm = sorted(m for s in cinfo["busy"].values() for m in s)
        span = f"{allm[0]//60:02d}:{allm[0]%60:02d}–{allm[-1]//60:02d}:{allm[-1]%60:02d}" if allm else "–"
        ws.cell(r, 1, n)
        ws.cell(r, 2, cinfo["name"])
        ws.cell(r, 3, ages_str(cinfo["ages"]) or "–")
        for i, d in enumerate(DAYS):
            ws.cell(r, 4 + i, f'=COUNTIFS({rng("B")},$B$3,{rng("K")},$B{r},{rng("D")},{get_column_letter(4+i)}${c_first-1})')
        ws.cell(r, 11, f"=SUM(D{r}:J{r})")
        ws.cell(r, 12, span)
        for col in range(1, 13):
            cell = ws.cell(r, col)
            cell.font = font(size=9, bold=(col == 2))
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if col in (2, 3) else "center")
        r += 1
    c_last = r - 1
    if coaches:
        ws.conditional_formatting.add(f"D{c_first}:J{c_last}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                                              end_type="max", end_color="8EC6D0"))
        dv3 = DataValidation(type="list", formula1=f"=$B${c_first}:$B${c_last}", allow_blank=True)
        ws.add_data_validation(dv3)
        if todo:
            dv3.add(f"M{first_item}:M{last_item}")
    else:
        ws.cell(r, 1, "Belum ada coach tercatat.").font = font(italic=True, color=C_MUTED)

    # Navigation links
    ws.cell(nav_row, 1, "Lompat ke").font = font(True, 9, C_MUTED)
    for i, (lab, row_) in enumerate([("Slot Map", slotmap_row), ("Daftar Diskusi", diskusi_row), ("Coach", coach_row)]):
        c = ws.cell(nav_row, 2 + i, f"→ {lab}")
        c.hyperlink = f"#'{center}'!A{row_}"
        c.font = Font(name=FONT, size=10, color=C_ACCENT, underline="single")
    ws.freeze_panes = "B1"


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("feed")
    ap.add_argument("out")
    ap.add_argument("--week", default="")
    ap.add_argument("--asof", default=date.today().isoformat())
    a = ap.parse_args()
    build(parse_feed(a.feed), a.out, a.week, a.asof)
