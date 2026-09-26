"""Bangun tab "Slot Map" baru untuk Exboard.

Semua rumus membaca tab 'Reguler Class Availability' (kolom A:L) lewat kolom Key (helper) di L,
format "KLM|- Lane 1 - Senin|08:15". Tiap sel menampilkan age group, terisi/kapasitas (%),
coach, dan assist coach sekaligus; warna sel mengikuti status keterisian.

Pemakaian:
    python build_slot_map_tab.py <feed.csv> <output.xlsx>

<feed.csv> dipakai untuk mengisi salinan tab Reguler Class Availability di file ini (hanya untuk
pratinjau). Setelah tab Slot Map disalin ke Exboard, rumusnya membaca tab asli di Exboard.
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RCA = "'Reguler Class Availability'"
CENTERS = ["KLM", "KWC", "PML", "TMP", "BTU", "HIB"]
DAYS = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
LANES = [1, 2, 3, 4, 5, 6]
ROWS_PER_LANE = {1: 16, 2: 14, 3: 14, 4: 14, 5: 14, 6: 14}   # jumlah jam terbanyak per lane + 2 cadangan
TIME_GRID = [f"{h:02d}:{m:02d}" for h in range(6, 21) for m in (0, 15, 30, 45)]  # 06:00–20:45, tiap 15 menit
AGES = [("Baby", "6–17 bulan"), ("Tiny", "18–36 bulan"), ("Little", "3–5 tahun"), ("Kids", "5–8+ tahun"), ("Star", "9–12 tahun")]

FONT = "Arial"
INK, MUTED, ACCENT = "1F2D33", "6B7C82", "0B6475"
STATUS = [  # (status, fill, label legenda)
    ("Belum terisi", "F8C9C0", "Belum terisi (kelas ada, 0 murid)"),
    ("Kosong", "FFE3A6", "Slot kosong (belum ada kelas)"),
    ("Sebagian", "EAF4F6", "Terisi < 60%"),
    ("Tinggi", "C4E1E6", "Terisi 60–99%"),
    ("Penuh", "BFE3C8", "Penuh"),
    ("Lebih", "E2D4F5", "Lebih kapasitas"),
    ("Tutup", "E4E4E4", "Tutup (kapasitas 0)"),
]
AGE_FILL = {"Baby": "F7DBE6", "Tiny": "FDE7C4", "Little": "D6EFDC", "Kids": "DBE4FB", "Star": "F1E2FA"}

# kolom
DISP0 = 2              # B..H = Senin..Minggu
IDX0 = 10              # J..P = nomor baris di Reguler Class Availability (helper)
ST0 = 18               # R..X = status (helper)
TZ, TF, TC = 26, 27, 28  # Z, AA, AB = daftar jam 15-menit, dipakai?, hitung kumulatif (helper)

thin = Side(style="thin", color="D5DDDF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(bold=False, size=10, color=INK, italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def fill(h):
    return PatternFill("solid", start_color=h, end_color=h)


def col(c):
    return get_column_letter(c)


def rc(letter):
    return f"{RCA}!${letter}:${letter}"


def num(expr):
    """Nilai angka dari sel Reguler Class Availability; kosong/teks non-angka = 0."""
    return f'IFERROR(--({expr}&""),0)'


def build_rca_copy(wb, feed):
    ws = wb.create_sheet("Reguler Class Availability")
    with open(feed, encoding="utf-8") as fh:
        for i, r in enumerate(csv.reader(fh)):
            r = (r + [""] * 12)[:12]
            if i > 0:
                for k in (5, 6):
                    try:
                        r[k] = int(r[k])
                    except ValueError:
                        pass
                r[3] = r[3].replace(" ", "").strip()
            ws.append([v if v != "" else None for v in r])
    ws["N1"] = "Salinan untuk pratinjau saja. Di Exboard, tab Slot Map membaca tab Reguler Class Availability yang asli."
    ws["N1"].font = font(True, 10, "B03A2E")


def build_slot_map(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 8
    for i in range(7):
        ws.column_dimensions[col(DISP0 + i)].width = 21
    ws.column_dimensions["I"].width = 2
    for c in list(range(IDX0, IDX0 + 7)) + [17] + list(range(ST0, ST0 + 7)) + [25, TZ, TF, TC]:
        ws.column_dimensions[col(c)].hidden = True

    # ---------- header ----------
    ws["A1"] = "SLOT MAP — JADWAL KELAS PER LANE"
    ws["A1"].font = font(True, 16, ACCENT)
    ws["A2"] = ("Baris = jam, kolom = hari. Tiap sel: age group · terisi/kapasitas (%) · coach · assist coach. "
                "Warna = status keterisian. Ganti center di B4.")
    ws["A2"].font = font(size=9, color=MUTED)
    ws["A4"] = "Center"
    ws["A4"].font = font(True, color=MUTED)
    ws["B4"] = "KLM"
    ws["B4"].font = font(True, 12)
    ws["B4"].fill = fill("FFF4B8")
    ws["B4"].border = BORDER
    ws["C4"] = "← pilih center"
    ws["C4"].font = font(size=9, color=MUTED, italic=True)
    dv = DataValidation(type="list", formula1='"' + ",".join(CENTERS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")

    # KPI (baris 6-7): kolom B..H
    st_all = f"${col(ST0)}$1:${col(ST0 + 6)}$1000"
    kpis = [
        ("Utilisasi center", f'=IFERROR(SUMIFS({rc("G")},{rc("A")},$B$4,{rc("F")},">0")/SUMIFS({rc("F")},{rc("A")},$B$4,{rc("F")},">0"),"–")', "0%", "DCEDF0"),
        ("Kursi terisi / kapasitas", f'=SUMIFS({rc("G")},{rc("A")},$B$4,{rc("F")},">0")&" / "&SUMIFS({rc("F")},{rc("A")},$B$4,{rc("F")},">0")', "@", "DCEDF0"),
        ("Belum terisi", f'=COUNTIF({st_all},"Belum terisi")', "0", "F8C9C0"),
        ("Slot kosong", f'=COUNTIF({st_all},"Kosong")', "0", "FFE3A6"),
        ("Penuh", f'=COUNTIF({st_all},"Penuh")', "0", "BFE3C8"),
        ("Lebih kapasitas", f'=COUNTIF({st_all},"Lebih")', "0", "E2D4F5"),
        ("Tutup", f'=COUNTIF({st_all},"Tutup")', "0", "E4E4E4"),
    ]
    for i, (lab, f, nf, bg) in enumerate(kpis):
        c = DISP0 + i
        a = ws.cell(6, c, lab)
        a.font = font(True, 9)
        b = ws.cell(7, c, f)
        b.font = font(True, 15)
        b.number_format = nf
        b.alignment = Alignment(horizontal="left")
        for x in (a, b):
            x.fill = fill(bg)
            x.border = BORDER
    ws.cell(6, 1, "Ringkasan").font = font(True, 9, MUTED)

    # Legenda status & age group
    ws.cell(9, 1, "Warna").font = font(True, 9, MUTED)
    for i, (st, bg, lab) in enumerate(STATUS):
        c = ws.cell(9, DISP0 + i, lab)
        c.fill = fill(bg)
        c.font = font(True, 8)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[9].height = 26
    ws.cell(10, 1, "Age group").font = font(True, 9, MUTED)
    for i, (a, rng) in enumerate(AGES):
        c = ws.cell(10, DISP0 + i, f"{a} · {rng}")
        c.fill = fill(AGE_FILL[a])
        c.font = font(True, 9)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")

    # ---------- helper daftar jam per lane (Z:AB) ----------
    helper_rows = {}
    h = 1
    ws.cell(h, TZ, "helper: jam 15-menit per lane — jangan diubah").font = font(size=8, color=MUTED)
    h += 1
    for lane in LANES:
        top = h
        for t in TIME_GRID:
            ws.cell(h, TZ, t)
            ws.cell(h, TF, f'=IF(COUNTIFS({rc("L")},$B$4&"|- Lane {lane} - *|"&{col(TZ)}{h})>0,1,0)')
            ws.cell(h, TC, f"={col(TF)}{h}" if h == top else f"={col(TF)}{h}+{col(TC)}{h-1}")
            h += 1
        helper_rows[lane] = (top, h - 1)

    # ---------- blok per lane ----------
    r = 12
    grid_blocks = []
    for lane in LANES:
        top_h, bot_h = helper_rows[lane]
        ws.cell(r, 1, f"LANE {lane}").font = font(True, 12, "FFFFFF")
        for c in range(1, 9):
            ws.cell(r, c).fill = fill(ACCENT)
        ws.cell(r, 3, (f'="Utilisasi lane: "&IFERROR(TEXT(SUMIFS({rc("G")},{rc("A")},$B$4,{rc("C")},"- Lane {lane} - *",{rc("F")},">0")'
                       f'/SUMIFS({rc("F")},{rc("A")},$B$4,{rc("C")},"- Lane {lane} - *",{rc("F")},">0"),"0%"),"–")'
                       f'&"   ·   Belum terisi: "&COUNTIF(${col(ST0)}${r+2}:${col(ST0+6)}${r+1+ROWS_PER_LANE[lane]},"Belum terisi")'
                       f'&"   ·   Slot kosong: "&COUNTIF(${col(ST0)}${r+2}:${col(ST0+6)}${r+1+ROWS_PER_LANE[lane]},"Kosong")'))
        ws.cell(r, 3).font = font(True, 10, "FFFFFF")
        r += 1
        hdr = r
        ws.cell(r, 1, "Jam").font = font(True, 9, MUTED)
        ws.cell(r, 1).border = BORDER
        for i, d in enumerate(DAYS):
            c = ws.cell(r, DISP0 + i, d)
            c.font = font(True, 10)
            c.fill = fill("F0F4F5")
            c.alignment = Alignment(horizontal="center")
            c.border = BORDER
        r += 1
        first = r
        for k in range(1, ROWS_PER_LANE[lane] + 1):
            jam = ws.cell(r, 1, f'=IFERROR(INDEX(${col(TZ)}${top_h}:${col(TZ)}${bot_h},MATCH({k},${col(TC)}${top_h}:${col(TC)}${bot_h},0)),"")')
            jam.font = font(True, 10)
            jam.alignment = Alignment(horizontal="center", vertical="top")
            jam.border = BORDER
            for i in range(7):
                dcol, icol, scol = col(DISP0 + i), col(IDX0 + i), col(ST0 + i)
                day_ref = f"{dcol}${hdr}"
                n = f"{icol}{r}"
                # nomor baris slot di Reguler Class Availability
                ws.cell(r, IDX0 + i, f'=IF($A{r}="","",IFERROR(MATCH($B$4&"|- Lane {lane} - "&{day_ref}&"|"&$A{r},{rc("L")},0),""))')
                age = f'INDEX({rc("B")},{n})&""'
                cap = num(f'INDEX({rc("F")},{n})')
                used = num(f'INDEX({rc("G")},{n})')
                coach = f'INDEX({rc("I")},{n})&""'
                asst = f'INDEX({rc("J")},{n})&""'
                # status (helper, dipakai pewarnaan)
                ws.cell(r, ST0 + i, (
                    f'=IF({n}="","",IF(AND({cap}=0,{used}=0),IF(OR({age}<>"",{coach}<>""),"Tutup","Tidak dibuka"),'
                    f'IF({used}>{cap},"Lebih",IF({used}={cap},"Penuh",IF({used}=0,IF({age}="","Kosong","Belum terisi"),'
                    f'IF({used}/{cap}>=0.6,"Tinggi","Sebagian"))))))'))
                # tampilan sel
                disp = ws.cell(r, DISP0 + i, (
                    f'=IF({n}="","",IF({scol}{r}="Tidak dibuka","–",'
                    f'IF({age}="","Kosong",PROPER({age}))&"  "&{used}&"/"&{cap}'
                    f'&IF({cap}>0,"  ("&TEXT({used}/{cap},"0%")&")","")'
                    f'&CHAR(10)&IF({coach}="","coach: –",PROPER({coach}))'
                    f'&IF({asst}="","",CHAR(10)&"asst: "&PROPER({asst}))))'))
                disp.font = font(size=9)
                disp.alignment = Alignment(wrap_text=True, vertical="top")
                disp.border = BORDER
            ws.row_dimensions[r].height = 40
            r += 1
        last = r - 1
        grid_blocks.append((first, last))
        # baris utilisasi per hari
        ws.cell(r, 1, "Util").font = font(True, 8, MUTED)
        for i in range(7):
            day_ref = f"{col(DISP0 + i)}${hdr}"
            c = ws.cell(r, DISP0 + i, (
                f'=IFERROR(TEXT(SUMIFS({rc("G")},{rc("A")},$B$4,{rc("C")},"- Lane {lane} - "&{day_ref},{rc("F")},">0")'
                f'/SUMIFS({rc("F")},{rc("A")},$B$4,{rc("C")},"- Lane {lane} - "&{day_ref},{rc("F")},">0"),"0%"),"–")'))
            c.font = font(True, 9, MUTED)
            c.fill = fill("F0F4F5")
            c.alignment = Alignment(horizontal="center")
            c.border = BORDER
        ws.cell(r, 1).fill = fill("F0F4F5")
        ws.cell(r, 1).border = BORDER
        r += 2

    # pewarnaan: sel tampilan B..H mengikuti status di R..X pada baris yang sama
    for first, last in grid_blocks:
        area = f"{col(DISP0)}{first}:{col(DISP0 + 6)}{last}"
        for st, bg, _ in STATUS:
            ws.conditional_formatting.add(area, FormulaRule(formula=[f'{col(ST0)}{first}="{st}"'], fill=fill(bg), stopIfTrue=True))
        ws.conditional_formatting.add(area, FormulaRule(formula=[f'{col(ST0)}{first}="Tidak dibuka"'], font=Font(name=FONT, color="B0B8BA"), stopIfTrue=True))

    ws.cell(r, 1, ("Catatan: sel kosong = jam itu tidak ada di jadwal lane tersebut. '–' = slot tidak dibuka (kapasitas 0, tanpa kelas). "
                   "Kolom J:AB tersembunyi berisi rumus bantu; jangan dihapus.")).font = font(size=8, color=MUTED, italic=True)
    ws.freeze_panes = "B1"


def main(feed, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Slot Map"
    build_slot_map(ws)
    build_rca_copy(wb, feed)
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
